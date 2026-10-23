"""Bulk inventory import from an Excel (.xlsx) or CSV file.

One row = one product, optionally with one stock batch. Products are matched by SKU: an existing SKU
is updated (blank cells leave the old value alone), a new SKU is created. A row with batch_no +
quantity also receives that stock as a new batch.

The whole file is validated first and imported in a single transaction: one bad row means nothing is
saved, and every problem is reported with its spreadsheet row number. Re-importing a batch that
already exists is an error rather than a top-up, so uploading the same file twice can't double stock.
"""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Batch, Category, Product

MAX_ROWS = 5000

# column key, header shown in the template, required for a NEW product, help text
COLUMNS = [
    ('sku', 'SKU / Barcode', True, 'Unique code; matches existing products'),
    ('name', 'Product name', True, 'Include strength and pack size'),
    ('category', 'Category', True, 'Created automatically if new'),
    ('purchase_price', 'Purchase price', True, 'Cost per unit you sell'),
    ('selling_price', 'Selling price', True, 'Price per unit you sell'),
    ('reorder_level', 'Reorder level', False, 'Default 10'),
    ('batch_no', 'Batch no', False, 'Leave blank to add/update the product only'),
    ('expiry_date', 'Expiry date', False, 'YYYY-MM-DD; required with a batch'),
    ('quantity', 'Quantity', False, 'Units received; required with a batch'),
    ('batch_cost', 'Batch unit cost', False, 'Defaults to purchase price'),
]
HEADER_ALIASES = {
    'sku': 'sku', 'sku / barcode': 'sku', 'barcode': 'sku', 'sku/barcode': 'sku',
    'name': 'name', 'product name': 'name', 'product': 'name',
    'category': 'category',
    'purchase price': 'purchase_price', 'purchase_price': 'purchase_price', 'cost price': 'purchase_price', 'cost': 'purchase_price',
    'selling price': 'selling_price', 'selling_price': 'selling_price', 'price': 'selling_price', 'sale price': 'selling_price',
    'reorder level': 'reorder_level', 'reorder_level': 'reorder_level',
    'batch no': 'batch_no', 'batch_no': 'batch_no', 'batch': 'batch_no', 'batch number': 'batch_no',
    'expiry date': 'expiry_date', 'expiry_date': 'expiry_date', 'expiry': 'expiry_date', 'exp date': 'expiry_date',
    'quantity': 'quantity', 'qty': 'quantity',
    'batch unit cost': 'batch_cost', 'batch_cost': 'batch_cost', 'batch cost': 'batch_cost',
}
DATE_FORMATS = ['%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y', '%d.%m.%Y', '%Y/%m/%d']


class ImportFileError(Exception):
    """The file as a whole can't be read (wrong type, no header, too many rows)."""


@dataclass
class RowResult:
    row: int
    sku: str
    name: str
    product_action: str = ''   # 'create' | 'update'
    batch_action: str = ''     # 'create' | ''
    quantity: int = 0
    errors: list = field(default_factory=list)
    clean: dict = field(default_factory=dict)

    def as_dict(self):
        return {
            'row': self.row, 'sku': self.sku, 'name': self.name, 'product_action': self.product_action,
            'batch_action': self.batch_action, 'quantity': self.quantity, 'errors': self.errors,
        }


# ---------- Reading the file ----------

def _read_rows(uploaded):
    name = (uploaded.name or '').lower()
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        try:
            sheet = load_workbook(uploaded, read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise ImportFileError('This file could not be opened as an Excel workbook.') from exc
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    if name.endswith('.csv'):
        text = uploaded.read().decode('utf-8-sig', errors='replace')
        return [row for row in csv.reader(io.StringIO(text))]
    raise ImportFileError('Upload an Excel (.xlsx) or CSV file. Older .xls files: open in Excel and "Save as" .xlsx.')


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value):
    if _blank(value):
        return ''
    # Excel stores long barcodes as numbers (8901234567890.0); keep them as plain digits.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _money(value, label, errors):
    if _blank(value):
        return None
    try:
        amount = Decimal(str(value).replace(',', '').strip()).quantize(Decimal('0.01'))
    except InvalidOperation:
        errors.append(f'{label} "{value}" is not a number.')
        return None
    if amount < 0:
        errors.append(f'{label} cannot be negative.')
        return None
    return amount


def _whole(value, label, errors):
    if _blank(value):
        return None
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        errors.append(f'{label} "{value}" is not a whole number.')
        return None
    if number != number.to_integral_value() or number < 0:
        errors.append(f'{label} must be a whole number of 0 or more.')
        return None
    return int(number)


def _date(value, errors):
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    errors.append(f'Expiry date "{text}" is not a date. Use YYYY-MM-DD, e.g. 2027-06-30.')
    return None


# ---------- Validation ----------

def analyse(uploaded):
    """Parse and validate the file without writing anything. Returns (results, summary)."""
    rows = _read_rows(uploaded)
    header_index = next((i for i, r in enumerate(rows) if any(not _blank(c) for c in r)), None)
    if header_index is None:
        raise ImportFileError('The file is empty.')

    header = [HEADER_ALIASES.get(_text(c).lower().rstrip(' *')) for c in rows[header_index]]
    missing = [label for key, label, required, _ in COLUMNS if required and key not in header]
    if missing:
        raise ImportFileError(f'Missing column(s): {", ".join(missing)}. Download the template to see the expected headers.')

    data_rows = [
        (header_index + 2 + offset, row) for offset, row in enumerate(rows[header_index + 1:])
        if any(not _blank(c) for c in row)
    ]
    if not data_rows:
        raise ImportFileError('The file has headers but no product rows.')
    if len(data_rows) > MAX_ROWS:
        raise ImportFileError(f'The file has {len(data_rows)} rows; the limit is {MAX_ROWS} per upload.')

    skus_in_file = set()
    existing_products = {p.sku.lower(): p for p in Product.objects.all()}
    existing_batches = set(
        (sku.lower(), no.lower()) for sku, no in Batch.objects.values_list('product__sku', 'batch_no')
    )
    batches_in_file = set()

    results = []
    for row_number, raw in data_rows:
        values = {key: raw[i] if i < len(raw) else None for i, key in enumerate(header) if key}
        errors = []
        sku = _text(values.get('sku'))
        name = _text(values.get('name'))
        result = RowResult(row=row_number, sku=sku, name=name)

        if not sku:
            errors.append('SKU / Barcode is required.')
        elif len(sku) > 64:
            errors.append('SKU / Barcode is longer than 64 characters.')
        elif sku.lower() in skus_in_file:
            errors.append(f'SKU {sku} appears more than once in this file. Use one row per product.')
        skus_in_file.add(sku.lower())

        product = existing_products.get(sku.lower()) if sku else None
        category = _text(values.get('category'))
        purchase = _money(values.get('purchase_price'), 'Purchase price', errors)
        selling = _money(values.get('selling_price'), 'Selling price', errors)
        reorder = _whole(values.get('reorder_level'), 'Reorder level', errors)

        if product is None:
            result.product_action = 'create'
            for label, value in [('Product name', name), ('Category', category),
                                 ('Purchase price', purchase), ('Selling price', selling)]:
                already_reported = any(e.startswith(label) for e in errors)
                if value in (None, '') and not already_reported:
                    errors.append(f'{label} is required for a new product.')
        else:
            result.product_action = 'update'
            result.name = name or product.name

        effective_purchase = purchase if purchase is not None else getattr(product, 'purchase_price', None)
        effective_selling = selling if selling is not None else getattr(product, 'selling_price', None)
        if effective_purchase is not None and effective_selling is not None and effective_selling < effective_purchase:
            errors.append('Selling price cannot be below purchase price.')
        if len(name) > 200:
            errors.append('Product name is longer than 200 characters.')
        if len(category) > 100:
            errors.append('Category is longer than 100 characters.')

        # Optional batch on the same row.
        batch_no = _text(values.get('batch_no'))
        expiry = _date(values.get('expiry_date'), errors)
        quantity = _whole(values.get('quantity'), 'Quantity', errors)
        batch_cost = _money(values.get('batch_cost'), 'Batch unit cost', errors)
        wants_batch = bool(batch_no) or expiry is not None or quantity is not None
        if wants_batch:
            if not batch_no:
                errors.append('Batch no is required when adding stock.')
            if expiry is None and not any('Expiry date' in e for e in errors):
                errors.append('Expiry date is required when adding stock.')
            if quantity is None and not any('Quantity' in e for e in errors):
                errors.append('Quantity is required when adding stock.')
            elif quantity == 0:
                errors.append('Quantity must be at least 1 when adding stock.')
            key = (sku.lower(), batch_no.lower())
            if batch_no and sku and key in existing_batches:
                errors.append(f'Batch {batch_no} already exists for this product. It was probably imported before.')
            elif batch_no and sku and key in batches_in_file:
                errors.append(f'Batch {batch_no} appears twice for this product in the file.')
            batches_in_file.add(key)
            result.batch_action = 'create'
            result.quantity = quantity or 0

        result.errors = errors
        result.clean = {
            'sku': sku, 'name': name, 'category': category, 'purchase_price': purchase,
            'selling_price': selling, 'reorder_level': reorder, 'batch_no': batch_no,
            'expiry_date': expiry, 'quantity': quantity,
            'batch_cost': batch_cost if batch_cost is not None else effective_purchase,
        }
        results.append(result)

    summary = {
        'rows': len(results),
        'products_to_create': sum(r.product_action == 'create' for r in results),
        'products_to_update': sum(r.product_action == 'update' for r in results),
        'batches_to_create': sum(r.batch_action == 'create' for r in results),
        'units_to_add': sum(r.quantity for r in results if r.batch_action == 'create'),
        'error_rows': sum(bool(r.errors) for r in results),
    }
    return results, summary


@transaction.atomic
def apply(results):
    """Write a fully validated import. Caller must ensure no row has errors."""
    categories = {c.name.lower(): c for c in Category.objects.all()}
    for result in results:
        data = result.clean
        category = None
        if data['category']:
            category = categories.get(data['category'].lower())
            if category is None:
                category = Category.objects.create(name=data['category'])
                categories[category.name.lower()] = category

        product = Product.objects.select_for_update().filter(sku__iexact=data['sku']).first()
        if product is None:
            product = Product.objects.create(
                sku=data['sku'], name=data['name'], category=category,
                purchase_price=data['purchase_price'], selling_price=data['selling_price'],
                reorder_level=data['reorder_level'] if data['reorder_level'] is not None else 10,
            )
        else:
            changes = {
                'name': data['name'] or None, 'category': category,
                'purchase_price': data['purchase_price'], 'selling_price': data['selling_price'],
                'reorder_level': data['reorder_level'],
            }
            for attr, value in changes.items():
                if value is not None:
                    setattr(product, attr, value)
            product.save()

        if result.batch_action == 'create':
            Batch.objects.create(
                product=product, batch_no=data['batch_no'], expiry_date=data['expiry_date'],
                quantity=data['quantity'], purchase_price=data['batch_cost'] or product.purchase_price,
            )


# ---------- Template ----------

def build_template():
    """An .xlsx with the expected headers, two example rows and a notes sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventory'
    header_fill = PatternFill('solid', fgColor='CCFBF1')
    for col, (key, label, required, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=f'{label} *' if required else label)
        cell.font = Font(bold=True, color='115E59')
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(label) + 6)
    ws.column_dimensions['B'].width = 34

    examples = [
        ['PANADOL-500', 'Panadol 500mg Capsule', 'Analgesics', 6.00, 8.00, 20, 'PN2401', date(2027, 6, 30), 50, 6.00],
        ['8901234567890', 'Amoxicillin 500mg Capsules (21)', 'Antibiotics', 2.60, 4.80, 15, 'AMX0925', date(2027, 3, 31), 40, None],
    ]
    for r, values in enumerate(examples, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, date):
                cell.number_format = 'yyyy-mm-dd'
            elif c in (4, 5, 10):
                cell.number_format = '0.00'
    # Keep SKUs as text so long barcodes don't turn into 8.9E+12.
    for r in range(2, MAX_ROWS + 2):
        ws.cell(row=r, column=1).number_format = '@'
    ws.freeze_panes = 'A2'

    notes = wb.create_sheet('How to fill')
    notes.column_dimensions['A'].width = 20
    notes.column_dimensions['B'].width = 70
    notes.append(['Column', 'What to enter'])
    notes['A1'].font = notes['B1'].font = Font(bold=True)
    for key, label, required, help_text in COLUMNS:
        notes.append([label + (' *' if required else ''), help_text])
    notes.append([])
    notes.append(['Rules', '* = required for a new product. For an existing SKU, blank cells keep the current value.'])
    notes.append(['', 'One row per product. Fill Batch no, Expiry date and Quantity to add stock in the same row.'])
    notes.append(['', 'Prices and quantities are per unit you sell (per capsule if you sell loose capsules).'])
    notes.append(['', 'Delete the example rows before uploading.'])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
