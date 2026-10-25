import io
from datetime import date
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from sales.tests.factories import make_batch, make_product, make_user

from .models import Batch, Category, Product

HEADERS = ['SKU / Barcode *', 'Product name *', 'Category *', 'Purchase price *', 'Selling price *',
           'Reorder level', 'Batch no', 'Expiry date', 'Quantity', 'Batch unit cost']
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def xlsx(*rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile('stock.xlsx', buffer.getvalue(), content_type=XLSX)


PANADOL = ['PANADOL-500', 'Panadol 500mg Capsule', 'Analgesics', 6, 8, 20, 'PN2401', date(2027, 6, 30), 50, None]


class InventoryImportTests(APITestCase):
    def setUp(self):
        self.client.force_authenticate(make_user())

    def upload(self, file, commit=False):
        return self.client.post('/api/inventory/import/', {'file': file, 'commit': 'true' if commit else 'false'},
                                format='multipart')

    def test_preview_writes_nothing(self):
        response = self.upload(xlsx(PANADOL))
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(response.data['committed'])
        self.assertEqual(response.data['summary']['products_to_create'], 1)
        self.assertEqual(response.data['summary']['units_to_add'], 50)
        self.assertFalse(Product.objects.exists())

    def test_commit_creates_category_product_and_batch(self):
        response = self.upload(xlsx(PANADOL), commit=True)
        self.assertEqual(response.status_code, 200, response.data)
        product = Product.objects.get(sku='PANADOL-500')
        self.assertEqual(product.category.name, 'Analgesics')
        self.assertEqual(product.selling_price, Decimal('8.00'))
        self.assertEqual(product.reorder_level, 20)
        batch = product.batches.get()
        self.assertEqual((batch.batch_no, batch.quantity, batch.expiry_date, batch.purchase_price),
                         ('PN2401', 50, date(2027, 6, 30), Decimal('6.00')))

    def test_existing_sku_is_updated_and_restocked_blank_cells_kept(self):
        product = make_product(sku='SKU1', selling_price='5.00', purchase_price='3.00', reorder_level=7)
        row = ['SKU1', '', '', '', '6.50', '', 'NEW1', '2028-01-31', 12, '']
        response = self.upload(xlsx(row), commit=True)
        self.assertEqual(response.status_code, 200, response.data)
        product.refresh_from_db()
        self.assertEqual(product.selling_price, Decimal('6.50'))
        self.assertEqual(product.purchase_price, Decimal('3.00'))  # blank → unchanged
        self.assertEqual(product.reorder_level, 7)
        self.assertEqual(product.batches.get(batch_no='NEW1').quantity, 12)

    def test_any_error_blocks_the_whole_import(self):
        bad = ['BAD', 'Bad row', 'Misc', 'abc', 5, '', '', '', '', '']
        preview = self.upload(xlsx(PANADOL, bad))
        self.assertEqual(preview.data['summary']['error_rows'], 1)
        self.assertIn('is not a number', preview.data['rows'][1]['errors'][0])
        self.assertEqual(preview.data['rows'][1]['row'], 3)  # spreadsheet row number

        commit = self.upload(xlsx(PANADOL, bad), commit=True)
        self.assertEqual(commit.status_code, 400)
        self.assertFalse(Product.objects.exists())
        self.assertFalse(Category.objects.exists())

    def test_reimporting_same_batch_is_rejected(self):
        self.upload(xlsx(PANADOL), commit=True)
        again = self.upload(xlsx(PANADOL))
        self.assertIn('already exists', again.data['rows'][0]['errors'][0])
        self.assertEqual(Batch.objects.get().quantity, 50)

    def test_batch_fields_must_come_together_and_selling_not_below_cost(self):
        rows = [
            ['A1', 'Thing', 'Misc', 5, 4, '', '', '', '', ''],              # selling < purchase
            ['A2', 'Thing 2', 'Misc', 1, 2, '', 'B1', '', 10, ''],          # missing expiry
        ]
        data = self.upload(xlsx(*rows)).data
        self.assertIn('below purchase price', data['rows'][0]['errors'][0])
        self.assertIn('Expiry date is required', ' '.join(data['rows'][1]['errors']))

    def test_numeric_barcode_and_csv_supported(self):
        csv_file = SimpleUploadedFile(
            'stock.csv',
            b'Barcode,Name,Category,Cost,Price,Batch,Expiry,Qty\n8901234567890,Cetirizine 10mg,Allergy,1.2,2.4,C1,30/06/2027,30\n',
            content_type='text/csv',
        )
        response = self.upload(csv_file, commit=True)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Product.objects.get().sku, '8901234567890')
        self.assertEqual(Batch.objects.get().expiry_date, date(2027, 6, 30))

    def test_missing_required_column_and_wrong_file_type(self):
        response = self.upload(xlsx(['x'], headers=['SKU']))
        self.assertEqual(response.status_code, 400)
        self.assertIn('Missing column', response.data['error']['message'])
        response = self.upload(SimpleUploadedFile('stock.pdf', b'%PDF', content_type='application/pdf'))
        self.assertIn('.xlsx', response.data['error']['message'])

    def test_template_download_is_valid_workbook(self):
        response = self.client.get('/api/inventory/import/template/')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.active['A1'].value, 'SKU / Barcode *')
        self.assertEqual(wb.active['B2'].value, 'Panadol 500mg Tablet')
        self.assertEqual(wb.active['G1'].value, 'Pieces per box')

    def test_existing_batch_from_other_product_does_not_clash(self):
        other = make_product(sku='OTHER')
        make_batch(other, 5, expires_in_days=100, batch_no='PN2401')
        response = self.upload(xlsx(PANADOL), commit=True)
        self.assertEqual(response.status_code, 200, response.data)


class ReceiveStockTests(APITestCase):
    def setUp(self):
        self.client.force_authenticate(make_user())
        from datetime import timedelta

        from django.utils import timezone
        self.expiry = (timezone.localdate() + timedelta(days=400)).isoformat()

    def receive(self, **body):
        return self.client.post('/api/inventory/receive/', {'quantity': 50, 'expiry_date': self.expiry, **body}, format='json')

    def test_new_medicine_with_total_cost_auto_codes(self):
        response = self.receive(
            new_product={'name': 'Panadol 500mg Tablet', 'category': 'Pain relief'}, total_cost='300', selling_price='8',
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(response.data['created_product'])
        product = Product.objects.get()
        self.assertTrue(product.sku.startswith('MED-'))
        self.assertEqual((product.purchase_price, product.selling_price), (Decimal('6.00'), Decimal('8.00')))
        batch = product.batches.get()
        self.assertEqual(batch.quantity, 50)
        self.assertTrue(batch.batch_no.startswith('EXP'))
        self.assertEqual(response.data['product']['stock'], 50)

    def test_existing_medicine_reuses_last_cost_and_price(self):
        product = make_product(selling_price='8.00', purchase_price='6.00')
        response = self.receive(product=product.pk, quantity=20)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertFalse(response.data['created_product'])
        self.assertEqual(product.batches.get().purchase_price, Decimal('6.00'))

    def test_two_deliveries_same_expiry_get_distinct_batch_numbers(self):
        product = make_product()
        first = self.receive(product=product.pk).data['batch']['batch_no']
        second = self.receive(product=product.pk).data['batch']['batch_no']
        self.assertNotEqual(first, second)

    def test_duplicate_name_is_blocked(self):
        make_product(name='Brufen 400mg Tablet')
        response = self.receive(new_product={'name': 'brufen 400MG tablet', 'category': 'X'}, total_cost='100', selling_price='5')
        self.assertEqual(response.status_code, 400)
        self.assertIn('already in your list', response.data['error']['message'])
        self.assertEqual(Product.objects.count(), 1)

    def test_selling_below_cost_and_past_expiry_rejected(self):
        loss = self.receive(new_product={'name': 'A', 'category': 'X'}, total_cost='500', selling_price='5')
        self.assertIn('lose money', loss.data['error']['message'])
        past = self.client.post('/api/inventory/receive/', {
            'product': make_product().pk, 'quantity': 1, 'expiry_date': '2020-01-31'}, format='json')
        self.assertIn('already passed', past.data['error']['message'])
        self.assertFalse(Batch.objects.exists())


class PaginationTests(APITestCase):
    """Every list endpoint pages, and pages never overlap or skip rows (needed for infinite scroll)."""

    def setUp(self):
        self.client.force_authenticate(make_user())
        # 25 products with the SAME price: ordering by price alone would be ambiguous.
        self.products = [make_product(name=f'Same {i:02d}', selling_price='10.00') for i in range(25)]

    def walk(self, url, **params):
        seen, page = [], 1
        while True:
            data = self.client.get(url, {**params, 'page': page, 'page_size': 10}).data
            self.assertEqual(data['page'], page)
            seen += [row['id'] for row in data['results']]
            if not data['next']:
                return data, seen
            page += 1

    def test_pages_cover_every_row_exactly_once_even_with_ties(self):
        for ordering in ['selling_price', '-selling_price', 'name', '']:
            data, seen = self.walk('/api/products/', ordering=ordering)
            self.assertEqual(len(seen), 25, ordering)
            self.assertEqual(len(set(seen)), 25, ordering)
            self.assertEqual(data['total_pages'], 3)

    def test_all_list_endpoints_are_paginated(self):
        for url in ['/api/products/', '/api/batches/', '/api/categories/', '/api/invoices/', '/api/returns/']:
            data = self.client.get(url, {'page_size': 5}).data
            for key in ('count', 'total_pages', 'page', 'page_size', 'next', 'previous', 'results'):
                self.assertIn(key, data, url)

    def test_page_size_is_capped(self):
        self.assertEqual(self.client.get('/api/products/', {'page_size': 5000}).data['page_size'], 200)

    def test_inventory_report_is_limited_but_counts_everything(self):
        for p in self.products:
            make_batch(p, 1, expires_in_days=10)
        data = self.client.get('/api/inventory/', {'limit': 5}).data
        self.assertEqual(len(data['expiring']), 5)
        self.assertEqual(data['expiring_count'], 25)
        self.assertEqual(len(data['low_stock']), 5)
        self.assertEqual(data['low_stock_count'], 25)
        self.assertEqual(self.client.get('/api/inventory/', {'limit': 0}).status_code, 400)
