"""Monthly report as an Excel workbook: the dashboard's figures plus the detail behind them."""

import io
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.db.models import DecimalField, ExpressionWrapper, F, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from inventory.models import Batch, Product
from sales.models import Invoice, InvoiceItem, Return, ReturnItem

from .views import build_dashboard

MONEY_FMT = '#,##0.00'
DATE_FMT = 'yyyy-mm-dd'
DATETIME_FMT = 'yyyy-mm-dd hh:mm'
TEAL = '0F766E'
HEADER_FILL = PatternFill('solid', fgColor='CCFBF1')
TOTAL_FILL = PatternFill('solid', fgColor='F0FDFA')
THIN = Side(style='thin', color='99F6E4')
DECIMAL = DecimalField(max_digits=14, decimal_places=2)


def _num(value):
    return float(Decimal(value or 0))


def _local(dt):
    # Excel has no time zones: write the pharmacy's local wall-clock time.
    return timezone.localtime(dt).replace(tzinfo=None)


def _table(ws, start_row, headers, rows, money_cols=(), date_cols=(), datetime_cols=(), widths=None, total_cols=()):
    """Write a header + rows block; optional bold totals row summing `total_cols`. Returns the next free row."""
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True, color='115E59')
        cell.fill = HEADER_FILL
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(horizontal='right' if col in money_cols or col in total_cols else 'left')
    r = start_row
    for r, values in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            if isinstance(value, str):
                # Customer names, reasons and product names are user input: never let "=..." become a formula.
                cell.data_type = 's'
            if col in money_cols:
                cell.number_format = MONEY_FMT
            elif col in date_cols:
                cell.number_format = DATE_FMT
            elif col in datetime_cols:
                cell.number_format = DATETIME_FMT
    if total_cols and rows:
        r += 1
        ws.cell(row=r, column=1, value='Total').font = Font(bold=True)
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).fill = TOTAL_FILL
        for col in total_cols:
            letter = get_column_letter(col)
            cell = ws.cell(row=r, column=col, value=f'=SUM({letter}{start_row + 1}:{letter}{r - 1})')
            cell.font = Font(bold=True)
            if col in money_cols:
                cell.number_format = MONEY_FMT
    if not rows:
        r += 1
        ws.cell(row=r, column=1, value='No records for this period.').font = Font(italic=True, color='7B8794')
    for col, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    return r + 2


def _title(ws, title, subtitle):
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=14, color=TEAL)
    ws['A2'] = subtitle
    ws['A2'].font = Font(color='52606D')
    return 4


def build_monthly_report(first, days_in_month, start, end):
    pharmacy_name = settings.PHARMACY_NAME
    data = build_dashboard(first, days_in_month, start, end)
    month_label = first.strftime('%B %Y')
    generated = timezone.localtime().strftime('%Y-%m-%d %H:%M')
    subtitle = f'{pharmacy_name} · {month_label} · generated {generated}'

    invoices = (
        Invoice.objects.filter(date__gte=start, date__lt=end)
        .select_related('created_by').order_by('date', 'id')
    )
    returns = (
        Return.objects.filter(date__gte=start, date__lt=end)
        .select_related('invoice', 'created_by').prefetch_related('items__invoice_item__product').order_by('date', 'id')
    )

    wb = Workbook()

    # ---------- Summary ----------
    ws = wb.active
    ws.title = 'Summary'
    row = _title(ws, f'Monthly report — {month_label}', subtitle)
    summary_rows = [
        ('Total sales', _num(data['total_sales'])),
        ('Returns (refunded)', _num(data['total_returns'])),
        ('Net revenue', _num(data['net_revenue'])),
        ('Profit', _num(data['total_profit'])),
        ('Invoices', data['invoice_count']),
        ('Average invoice', _num(data['total_sales']) / data['invoice_count'] if data['invoice_count'] else 0),
    ]
    row = _table(ws, row, ['Measure', 'Value'], summary_rows, money_cols=(2,), widths=[28, 18])
    for r in range(row - 2 - len(summary_rows) + 1, row - 1):
        if ws.cell(row=r, column=1).value == 'Invoices':
            ws.cell(row=r, column=2).number_format = '0'
    ws.cell(row=row, column=1, value='Stock alerts (as of today)').font = Font(bold=True, color=TEAL)
    _table(ws, row + 1, ['Alert', 'Count'], [
        ('Low-stock products', data['low_stock_count']),
        (f'Batches expiring within {data["expiry_warning_days"]} days', data['expiring_count']),
        ('Expired batches still in stock', data['expired_count']),
    ], widths=[34, 18])

    # ---------- Daily ----------
    ws = wb.create_sheet('Daily')
    row = _title(ws, 'Daily sales', subtitle)
    _table(
        ws, row, ['Date', 'Sales', 'Returns', 'Net sales', 'Profit'],
        [(date.fromisoformat(d['date']), _num(d['sales']), _num(d['returns']),
          _num(d['net_sales']), _num(d['profit'])) for d in data['daily']],
        money_cols=(2, 3, 4, 5), date_cols=(1,), total_cols=(2, 3, 4, 5), widths=[14, 14, 14, 14, 14],
    )

    # ---------- Products sold ----------
    ws = wb.create_sheet('Products sold')
    row = _title(ws, 'Sales by product', subtitle)
    sold = {
        r['product_id']: r for r in InvoiceItem.objects.filter(invoice__date__gte=start, invoice__date__lt=end)
        .values('product_id', 'product__name', 'product__sku', 'product__category__name')
        .annotate(
            qty=Sum('quantity'), revenue=Sum('line_total'),
            cost=Sum(ExpressionWrapper(F('unit_cost') * F('quantity'), output_field=DECIMAL)),
        )
    }
    returned = {
        r['invoice_item__product_id']: r for r in ReturnItem.objects.filter(return_obj__date__gte=start, return_obj__date__lt=end)
        .values('invoice_item__product_id', 'invoice_item__product__name', 'invoice_item__product__sku',
                'invoice_item__product__category__name')
        .annotate(
            qty=Sum('quantity'), refund=Sum('refund_amount'),
            lost_profit=Sum(ExpressionWrapper(
                (F('invoice_item__unit_price') - F('invoice_item__unit_cost')) * F('quantity'), output_field=DECIMAL)),
        )
    }
    product_rows = []
    # Union: a product returned this month but sold last month still counts against this month's profit.
    for pid in sold.keys() | returned.keys():
        s, ret = sold.get(pid, {}), returned.get(pid, {})
        name = s.get('product__name') or ret.get('invoice_item__product__name')
        sku = s.get('product__sku') or ret.get('invoice_item__product__sku')
        category = s.get('product__category__name') or ret.get('invoice_item__product__category__name')
        gross_profit = (s.get('revenue') or 0) - (s.get('cost') or 0)
        product_rows.append((
            name, sku, category, s.get('qty') or 0, ret.get('qty') or 0,
            _num(s.get('revenue')), _num(ret.get('refund')), _num(s.get('cost')),
            _num(gross_profit - (ret.get('lost_profit') or 0)),
        ))
    product_rows.sort(key=lambda r: r[5], reverse=True)
    _table(
        ws, row, ['Product', 'SKU', 'Category', 'Qty sold', 'Qty returned', 'Revenue', 'Refunded', 'Cost', 'Profit'],
        product_rows, money_cols=(6, 7, 8, 9), total_cols=(4, 5, 6, 7, 8, 9),
        widths=[36, 18, 18, 10, 13, 14, 14, 14, 14],
    )

    # ---------- Invoices ----------
    ws = wb.create_sheet('Invoices')
    row = _title(ws, 'Invoices', subtitle)
    _table(
        ws, row, ['Invoice', 'Date', 'Customer', 'Cashier', 'Payment', 'Subtotal', 'Discount', 'Total', 'Status'],
        [(i.number, _local(i.date), i.customer_name or 'Walk-in', i.created_by.get_full_name() or i.created_by.username,
          i.get_payment_method_display(), _num(i.subtotal), _num(i.discount), _num(i.total), i.get_status_display())
         for i in invoices],
        money_cols=(6, 7, 8), datetime_cols=(2,), total_cols=(6, 7, 8),
        widths=[14, 18, 20, 18, 10, 13, 12, 13, 18],
    )

    # ---------- Returns ----------
    ws = wb.create_sheet('Returns')
    row = _title(ws, 'Returns', subtitle)
    _table(
        ws, row, ['Return', 'Date', 'Invoice', 'Items', 'Reason', 'Processed by', 'Refund'],
        [(r.number, _local(r.date), r.invoice.number,
          ', '.join(f'{it.quantity} x {it.invoice_item.product.name}' for it in r.items.all()),
          r.reason, r.created_by.get_full_name() or r.created_by.username, _num(r.refund_amount))
         for r in returns],
        money_cols=(7,), datetime_cols=(2,), total_cols=(7,), widths=[13, 18, 14, 40, 28, 18, 12],
    )

    # ---------- Stock alerts ----------
    ws = wb.create_sheet('Stock alerts')
    row = _title(ws, 'Stock alerts (as of today)', subtitle)
    ws.cell(row=row, column=1, value='Low stock').font = Font(bold=True, color=TEAL)
    row = _table(
        ws, row + 1, ['Product', 'SKU', 'Sellable stock', 'Reorder level'],
        [(p.name, p.sku, p.sellable_stock, p.reorder_level)
         for p in Product.objects.low_stock().order_by('sellable_stock', 'name')],
        widths=[36, 18, 15, 14],
    )
    today = timezone.localdate()
    batches = Batch.objects.select_related('product').in_stock()
    for label, qs in [
        (f'Expiring within {settings.EXPIRY_WARNING_DAYS} days', batches.expiring_within(settings.EXPIRY_WARNING_DAYS, today)),
        ('Expired (still in stock)', batches.expired(today)),
    ]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=TEAL)
        row = _table(
            ws, row + 1, ['Product', 'Batch', 'Expiry date', 'Quantity', 'Cost value'],
            [(b.product.name, b.batch_no, b.expiry_date, b.quantity, _num(b.purchase_price * b.quantity))
             for b in qs.fefo()],
            money_cols=(5,), date_cols=(3,), total_cols=(4, 5), widths=[36, 18, 15, 14, 14],
        )

    for sheet in wb.worksheets:
        sheet.freeze_panes = None
        sheet.sheet_view.showGridLines = False

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
